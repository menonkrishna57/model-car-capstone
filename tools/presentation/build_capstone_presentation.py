"""
Builds the base capstone presentation (.pptx) for the Model-Car Capstone.
Dashed boxes marked "SCREENSHOT" are placeholders to be replaced with real
captures (benchmarks, detections, car photos). Table cells marked "__" are
to be filled with numbers measured on the Raspberry Pi 5.

Run:  python build_capstone_presentation.py
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
from lxml import etree

OUT_DIR = r"e:\Capstone\model-car-capstone\docs\reports"
OUT_FILE = "capstone_presentation_base.pptx"

# ---------------- PALETTE ----------------
BG = RGBColor(0x0B, 0x12, 0x20)
CARD = RGBColor(0x13, 0x1C, 0x2E)
CARD2 = RGBColor(0x1B, 0x26, 0x3F)
CYAN = RGBColor(0x22, 0xD3, 0xEE)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
GREEN = RGBColor(0x34, 0xD3, 0x99)
RED = RGBColor(0xF8, 0x71, 0x71)
VIOLET = RGBColor(0xA7, 0x8B, 0xFA)
TEXT = RGBColor(0xE5, 0xE7, 0xEB)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
FONT = "Segoe UI"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
slide_no = [0]


# ---------------- HELPERS ----------------
def new_slide(title=None, subtitle=None, notes=None):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    slide_no[0] += 1
    if title:
        text(s, 0.6, 0.35, 12, 0.7, title, 30, TEXT, bold=True)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.05), Inches(1.2), Inches(0.06))
        fill(bar, CYAN)
        if subtitle:
            text(s, 2.0, 0.93, 10.7, 0.35, subtitle, 13, MUTED)
    if slide_no[0] > 1:
        text(s, 0.6, 7.05, 8, 0.3, "Model-Car Capstone  |  NMIMS MPSTME", 10, MUTED)
        text(s, 12.0, 7.05, 0.8, 0.3, str(slide_no[0]), 10, MUTED, align=PP_ALIGN.RIGHT)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def fill(shape, color, line=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1.25)
    shape.shadow.inherit = False


def text(s, x, y, w, h, content, size=14, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, italic=False):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    lines = content if isinstance(content, list) else [content]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = ln
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
    return tb


def bullets(s, x, y, w, h, items, size=15, color=TEXT, bullet_color=CYAN, spacing=6):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(spacing)
        if isinstance(item, tuple):
            head, body = item
        else:
            head, body = None, item
        b = p.add_run()
        b.text = "▸ "
        b.font.name = FONT
        b.font.size = Pt(size)
        b.font.color.rgb = bullet_color
        if head:
            r = p.add_run()
            r.text = head + " "
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = True
            r.font.color.rgb = color
        r = p.add_run()
        r.text = body
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.color.rgb = color
    return tb


def card(s, x, y, w, h, title, lines, accent=CYAN, title_size=16, body_size=13):
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.adjustments[0] = 0.06
    fill(box, CARD)
    strip = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y + 0.15), Inches(0.07), Inches(h - 0.3))
    fill(strip, accent)
    text(s, x + 0.25, y + 0.15, w - 0.4, 0.45, title, title_size, accent, bold=True)
    bullets(s, x + 0.25, y + 0.62, w - 0.4, h - 0.7, lines, size=body_size, bullet_color=accent, spacing=4)


def node(s, x, y, w, h, title, sub=None, color=CYAN, title_size=12, sub_size=9.5):
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.adjustments[0] = 0.12
    fill(box, CARD2, line=color)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.06)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    r.font.name = FONT
    r.font.size = Pt(title_size)
    r.font.bold = True
    r.font.color.rgb = color
    if sub:
        for ln in sub.split("\n"):
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            r2 = p2.add_run()
            r2.text = ln
            r2.font.name = FONT
            r2.font.size = Pt(sub_size)
            r2.font.color.rgb = TEXT
    return box


def arrow(s, x1, y1, x2, y2, color=MUTED, width=1.75):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return c


def placeholder(s, x, y, w, h, label, hint=None):
    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0x0F, 0x18, 0x2A)
    box.line.color.rgb = AMBER
    box.line.width = Pt(1.5)
    box.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "📷  SCREENSHOT"
    r.font.name = FONT
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = AMBER
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = label
    r2.font.name = FONT
    r2.font.size = Pt(12)
    r2.font.color.rgb = TEXT
    if hint:
        p3 = tf.add_paragraph()
        p3.alignment = PP_ALIGN.CENTER
        r3 = p3.add_run()
        r3.text = hint
        r3.font.name = FONT
        r3.font.size = Pt(9.5)
        r3.font.italic = True
        r3.font.color.rgb = MUTED
    return box


def table(s, x, y, w, h, rows, col_widths=None, header_color=CYAN, size=11, highlight_row=None):
    nr, nc = len(rows), len(rows[0])
    gt = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = gt.table
    # strip default table style banding
    tblPr = gt._element.graphic.graphicData.tbl.tblPr
    tblPr.set("bandRow", "0")
    tblPr.set("firstRow", "0")
    if col_widths:
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = Inches(cw)
    for r in range(nr):
        for c in range(nc):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = CARD2
            elif highlight_row is not None and r == highlight_row:
                cell.fill.fore_color.rgb = RGBColor(0x0E, 0x3A, 0x45)
            else:
                cell.fill.fore_color.rgb = CARD if r % 2 else RGBColor(0x10, 0x18, 0x28)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if c > 0 else PP_ALIGN.LEFT
            run = p.add_run()
            run.text = str(rows[r][c])
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.bold = (r == 0) or (highlight_row is not None and r == highlight_row and c == 0)
            run.font.color.rgb = header_color if r == 0 else (AMBER if rows[r][c] == "__" else TEXT)
    return tbl


# =====================================================================
# SLIDE 1 — TITLE
# =====================================================================
s = new_slide(notes="Introduce the team and the one-line pitch: a low-cost two-layer autonomous car "
                    "that does person following, obstacle avoidance and pothole detection.")
glow = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.3), Inches(-1.5), Inches(7), Inches(7))
fill(glow, RGBColor(0x0E, 0x2A, 0x3A))
text(s, 0.8, 1.2, 9, 0.4, "B.TECH CAPSTONE PROJECT  •  NMIMS MPSTME", 13, CYAN, bold=True)
text(s, 0.8, 1.75, 11.5, 2.0,
     ["Autonomous Follow-Me & Obstacle-Avoidant", "Robotic Vehicle"], 40, TEXT, bold=True)
text(s, 0.8, 3.55, 11, 0.9,
     "Two-Layer Architecture (Raspberry Pi 5 + Arduino Uno) with NMS-Free YOLOv10n on Tencent NCNN, "
     "Pothole Detection and Hard Real-Time Reflex Control", 16, MUTED)
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(4.65), Inches(1.6), Inches(0.06))
fill(bar, CYAN)
text(s, 0.8, 4.9, 8, 0.4, "Krishna Menon   •   Aditya Rana   •   Parth Surve", 18, TEXT, bold=True)
text(s, 0.8, 5.35, 8, 0.4, "Guide: ____________________", 14, MUTED)
placeholder(s, 9.2, 4.2, 3.4, 2.5, "Photo of the assembled car", "Front 3/4 view, camera visible")

# =====================================================================
# SLIDE 2 — PROBLEM & RESEARCH GAPS
# =====================================================================
s = new_slide("Problem Statement & Research Gaps",
              "Why a single Linux board driving motors directly is not enough",
              notes="Gap 1: Linux is not real-time; PWM and braking latency suffer under heavy inference. "
                    "Gap 2: ultrasonic sensors cannot see potholes. Gap 3: commercial kits are expensive in India.")
card(s, 0.6, 1.5, 3.9, 4.0, "01  Timing & Safety", [
    "Linux CFS scheduler gives no real-time guarantees",
    "Software PWM jitters under heavy neural inference",
    "Camera → decision → motor loop is too slow for last-moment braking",
], accent=RED)
card(s, 4.72, 1.5, 3.9, 4.0, "02  Negative Obstacles", [
    "Ultrasonic waves travel horizontally",
    "Potholes / drop-offs give no echo → reads as \"path clear\"",
    "Indian roads make pothole awareness essential",
], accent=AMBER)
card(s, 8.84, 1.5, 3.9, 4.0, "03  Cost & Access", [
    "Commercial AI car kits are imported & costly",
    "Need a modular build from locally available parts",
    "Re-use of college lab assets (Pi 5, Arduino, camera)",
], accent=GREEN)
box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(5.75), Inches(12.14), Inches(1.0))
fill(box, CARD2, line=CYAN)
text(s, 0.9, 5.85, 11.6, 0.8,
     "Objective: decouple perception (Pi 5) from actuation (Arduino) so that vision can be heavy "
     "while braking and motor timing stay deterministic — and add camera-based pothole detection.",
     14, TEXT, anchor=MSO_ANCHOR.MIDDLE)

# =====================================================================
# SLIDE 3 — SYSTEM ARCHITECTURE
# =====================================================================
s = new_slide("System Architecture — The Two-Layer Split",
              "Cognitive brain + reflexive spine, connected over USB serial",
              notes="Pi 5 runs all vision and decision making. Arduino owns PWM, sonar timing and the reflex. "
                    "If the Pi crashes, the Arduino watchdog stops the motors.")
node(s, 0.6, 1.5, 4.6, 3.3, "Raspberry Pi 5 (8 GB)  —  High-Level",
     "Quad Cortex-A76 @ 2.4 GHz\n\n• CSI camera ingest (640×480)\n• YOLOv10n on Tencent NCNN\n"
     "• Pothole engine (RDD2022 YOLOv10n + CV)\n• KCF tracker + HSV re-ID\n• PID steering & decision arbiter",
     color=CYAN, title_size=15, sub_size=12)
node(s, 7.5, 1.5, 4.6, 3.3, "Arduino Uno R3  —  Real-Time",
     "ATmega328P @ 16 MHz\n\n• Hardware Timer0 PWM (D5 / D6)\n• 2× HC-SR04P sonar, 2× IR\n"
     "• 3-tier priority state machine\n• Local reflex avoidance (< 20 cm)\n• Watchdog failsafe",
     color=GREEN, title_size=15, sub_size=12)
arrow(s, 5.3, 2.75, 7.4, 2.75, CYAN, 2.5)
arrow(s, 7.4, 3.45, 5.3, 3.45, GREEN, 2.5)
text(s, 5.25, 2.2, 2.2, 0.5, "DRIVE / TURN / PI_STOP", 10, CYAN, align=PP_ALIGN.CENTER)
text(s, 5.25, 3.55, 2.2, 0.5, "[TELEM] / [ACK]", 10, GREEN, align=PP_ALIGN.CENTER)
text(s, 5.25, 2.95, 2.2, 0.4, "USB 115200 8N1", 11, TEXT, bold=True, align=PP_ALIGN.CENTER)
node(s, 0.6, 5.1, 4.6, 1.15, "Logic Rail (isolated)", "USB-C power bank 5.1 V / 5 A → Pi 5", color=VIOLET, sub_size=11)
node(s, 7.5, 5.1, 4.6, 1.15, "Motor Rail", "2S 18650 (7.4 V) → buck 5 V → Arduino / L298N → 4× TT motors",
     color=AMBER, sub_size=11)

# =====================================================================
# SLIDE 4 — WHY YOLOv10n
# =====================================================================
s = new_slide("Perception Model — Why YOLOv10n?",
              "End-to-end detection without Non-Maximum Suppression",
              notes="Older YOLOs output many overlapping boxes and need NMS on the CPU. YOLOv10 trains with "
                    "dual label assignment and uses the one-to-one head at inference, so NMS is removed.")
text(s, 0.6, 1.45, 4.0, 0.4, "Traditional YOLO (v5 / v8)", 15, RED, bold=True)
ys = [1.95, 2.95, 3.95, 4.95]
labels_l = [("Backbone + Neck", None), ("One-to-many head", "many overlapping candidate boxes"),
            ("NMS post-processing", "IoU sort & suppress on CPU"), ("Final detections", None)]
for (t, sub), y in zip(labels_l, ys):
    node(s, 0.6, y, 3.6, 0.75, t, sub, color=RED if "NMS" in t else MUTED)
for y in ys[:-1]:
    arrow(s, 2.4, y + 0.75, 2.4, y + 1.0)

text(s, 4.7, 1.45, 4.0, 0.4, "YOLOv10 (NMS-free)", 15, GREEN, bold=True)
labels_r = [("Backbone + Neck", None), ("One-to-one head", "one box per object"),
            ("Final detections", "no NMS step")]
ys_r = [1.95, 3.45, 4.95]
for (t, sub), y in zip(labels_r, ys_r):
    node(s, 4.7, y, 3.6, 0.75, t, sub, color=GREEN if t != "Backbone + Neck" else MUTED)
arrow(s, 6.5, 2.7, 6.5, 3.45)
arrow(s, 6.5, 4.2, 6.5, 4.95)

card(s, 8.9, 1.45, 3.85, 5.3, "YOLOv10n at a glance", [
    ("Params:", "2.3 M"),
    ("Compute:", "6.7 GFLOPs (640 input)"),
    ("Trained on:", "COCO, 80 classes"),
    ("Classes used:", "person, bicycle, car, motorcycle, bus, truck"),
    ("Training trick:", "dual label assignment (one-to-many + one-to-one)"),
    ("Inference:", "one-to-one head only → no NMS"),
    ("Source:", "Wang et al., arXiv:2405.14458 (2024)"),
], accent=CYAN, body_size=12)

# =====================================================================
# SLIDE 5 — WHY NCNN
# =====================================================================
s = new_slide("Inference Engine — Why Tencent NCNN on the Pi 5 CPU?",
              "Making a CPU-only board run real-time detection",
              notes="Explain each pillar briefly, then point at the benchmark screenshot comparing "
                    "PyTorch, ONNX, NCNN and OpenVINO on the Pi 5.")
pillars = [
    ("ARM NEON SIMD", "128-bit vector instructions on Cortex-A76 process several values per instruction", CYAN),
    ("FP16 weights", "Half-precision export halves model size and memory bandwidth", GREEN),
    ("Operator fusion", "Conv + BatchNorm + activation merged into a single pass", AMBER),
    ("Light runtime", "No Python/PyTorch overhead in the inner loop; pre-allocated buffers", VIOLET),
]
for i, (t, d, col) in enumerate(pillars):
    y = 1.5 + i * 1.3
    node(s, 0.6, y, 1.9, 1.1, t, None, color=col, title_size=13)
    text(s, 2.65, y + 0.1, 3.6, 1.0, d, 12.5, TEXT, anchor=MSO_ANCHOR.MIDDLE)
placeholder(s, 6.5, 1.5, 6.25, 3.6, "benchmark.py output on Raspberry Pi 5",
            "PyTorch FP32 vs ONNX vs NCNN FP16 vs OpenVINO FP16 — latency / FPS table")
text(s, 6.5, 5.25, 6.25, 1.5,
     ["Export pipeline:", "yolov10n.pt  →  model.export(format=\"ncnn\", imgsz=320, half=True)  →  yolov10n_ncnn_model/"],
     12, MUTED)

# =====================================================================
# SLIDE 6 — COMPLETE SOFTWARE FLOWCHART
# =====================================================================
s = new_slide("Complete Software Pipeline",
              "Camera frame → perception → decision → serial command → reflex firmware → motors",
              notes="Walk left-to-right across the Pi 5 row, down through the serial link, then right-to-left "
                    "across the Arduino row. Emphasise that sonar feeds the Arduino directly, bypassing Linux.")
# region frames
pi_reg = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), Inches(1.35), Inches(12.55), Inches(2.65))
pi_reg.adjustments[0] = 0.04
fill(pi_reg, RGBColor(0x0E, 0x17, 0x28), line=RGBColor(0x1E, 0x3A, 0x5F))
text(s, 0.55, 1.38, 3, 0.3, "RASPBERRY PI 5  (Python)", 10, CYAN, bold=True)
ar_reg = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), Inches(4.85), Inches(12.55), Inches(2.1))
ar_reg.adjustments[0] = 0.05
fill(ar_reg, RGBColor(0x0E, 0x17, 0x28), line=RGBColor(0x1E, 0x4A, 0x3A))
text(s, 0.55, 4.88, 6, 0.3, "ARDUINO UNO  (C++ firmware, 100 Hz loop)", 10, GREEN, bold=True)

node(s, 0.6, 2.2, 1.55, 0.95, "Camera", "640×480 @ 30 FPS")
node(s, 2.45, 2.2, 1.55, 0.95, "Pre-process", "resize / ROI crop")
node(s, 4.3, 1.7, 2.0, 0.85, "YOLOv10n · NCNN", "people, vehicles", color=CYAN)
node(s, 4.3, 2.95, 2.0, 0.85, "Pothole engine", "RDD2022 YOLOv10n + CV", color=AMBER)
node(s, 6.6, 2.2, 1.8, 0.95, "Zone scorer", "Left / Centre / Right", color=VIOLET)
node(s, 8.7, 2.2, 1.6, 0.95, "Decision arbiter", None, color=VIOLET)
node(s, 10.6, 1.6, 2.2, 0.6, "PI_STOP", "emergency latch", color=RED, sub_size=8.5)
node(s, 10.6, 2.35, 2.2, 0.6, "PI_AVOID L/R", "directional swerve", color=AMBER, sub_size=8.5)
node(s, 10.6, 3.1, 2.2, 0.75, "Follow-Me", "KCF + HSV re-ID → PID", color=GREEN, sub_size=8.5)
arrow(s, 2.15, 2.67, 2.45, 2.67)
arrow(s, 4.0, 2.5, 4.3, 2.12)
arrow(s, 4.0, 2.85, 4.3, 3.37)
arrow(s, 6.3, 2.12, 6.6, 2.5)
arrow(s, 6.3, 3.37, 6.6, 2.85)
arrow(s, 8.4, 2.67, 8.7, 2.67)
arrow(s, 10.3, 2.5, 10.6, 1.9)
arrow(s, 10.3, 2.67, 10.6, 2.65)
arrow(s, 10.3, 2.85, 10.6, 3.47)

serial = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), Inches(4.13), Inches(12.55), Inches(0.55))
fill(serial, CARD2, line=CYAN)
tf = serial.text_frame
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "USB-Serial 115200 8N1  —  ASCII commands: DRIVE · TURN · REV · STOP · PI_STOP · PI_GO · PI_AVOID      ⇅      [TELEM] / [ACK] replies"
r.font.name = FONT
r.font.size = Pt(11)
r.font.color.rgb = TEXT
arrow(s, 11.7, 3.85, 11.7, 4.13, CYAN)
arrow(s, 11.7, 4.68, 11.7, 5.15, CYAN)

node(s, 10.6, 5.15, 2.2, 0.95, "Tier 1", "PI_STOP latch / watchdog", color=RED)
node(s, 8.1, 5.15, 2.2, 0.95, "Tier 2", "sonar < 20 cm → reverse, turn, resume", color=AMBER)
node(s, 5.6, 5.15, 2.2, 0.95, "Tier 3", "execute DRIVE / TURN", color=GREEN)
node(s, 3.1, 5.15, 2.2, 0.95, "L298N driver", "PWM D5 / D6, dir D8–D11")
node(s, 0.6, 5.15, 2.2, 0.95, "4× TT motors", "skid-steer")
node(s, 8.1, 6.25, 2.2, 0.55, "2× HC-SR04P + 2× IR", None, color=MUTED, title_size=10)
arrow(s, 10.6, 5.62, 10.3, 5.62)
arrow(s, 8.1, 5.62, 7.8, 5.62)
arrow(s, 5.6, 5.62, 5.3, 5.62)
arrow(s, 3.1, 5.62, 2.8, 5.62)
arrow(s, 9.2, 6.25, 9.2, 6.1, MUTED)

# =====================================================================
# SLIDE 7 — POTHOLE ENGINE
# =====================================================================
s = new_slide("Negative Obstacles — Dual-Path Pothole Engine",
              "What ultrasonic sensors cannot see, the camera must",
              notes="Path A is the trained deep model; Path B is a fast classical fallback. Both feed the same "
                    "zone scorer, with potholes weighted higher than normal obstacles.")
text(s, 0.6, 1.45, 6, 0.4, "Path A — Deep model", 15, AMBER, bold=True)
bullets(s, 0.6, 1.85, 6, 1.6, [
    "Custom YOLOv10n fine-tuned on RDD2022 (India subset, D40 pothole class)",
    "50 training epochs  •  weights: best.pt (YOLOv10n)",
    "Validation: Precision 54.0 %, mAP@50 25.2 %",
], size=13, bullet_color=AMBER)
text(s, 0.6, 3.35, 6, 0.4, "Path B — Classical CV fallback", 15, CYAN, bold=True)
steps = ["ROI\nbottom 55 %", "Gaussian\n5×5", "Laplacian\nedges", "Morphology\n13×13 ellipse", "Contour\nfilters"]
for i, st in enumerate(steps):
    x = 0.6 + i * 1.22
    node(s, x, 3.85, 1.05, 0.95, st.split("\n")[0], st.split("\n")[1], color=CYAN, title_size=10.5, sub_size=9)
    if i < len(steps) - 1:
        arrow(s, x + 1.05, 4.32, x + 1.22, 4.32)
text(s, 0.6, 5.0, 6.1, 0.9,
     "Contour filters: 900 px < area < 25 % of ROI, aspect ratio 0.25–5.0",
     11.5, MUTED)
box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(5.65), Inches(6.1), Inches(1.15))
fill(box, CARD2, line=VIOLET)
text(s, 0.8, 5.7, 5.8, 1.05,
     ["Zone danger score", "Danger = conf × W × (0.4 + min(A_box / A_frame, 0.5))",
      "W = 1.0 normal obstacle  •  W = 1.2–2.0 pothole"], 12, TEXT)
placeholder(s, 7.0, 1.45, 5.75, 2.55, "Pothole detection output frame", "Corridor overlay + swerve decision")
placeholder(s, 7.0, 4.2, 5.75, 2.6, "Training curves (results.png)", "runs_rdd/custom_pothole_model — loss & mAP vs epoch")

# =====================================================================
# SLIDE 8 — FOLLOW-ME TRACKING
# =====================================================================
s = new_slide("Follow-Me Tracking & Re-Identification",
              "Lock onto one person, keep following, recover after occlusion",
              notes="KCF keeps the box between YOLO frames. If the target is lost, the car stops, YOLO finds "
                    "candidates, and the HSV torso histogram picks the right person again.")
flow = [("Target lock", "YOLO person box", CYAN), ("KCF tracker", "fast frame-to-frame", CYAN),
        ("Lost?", "occlusion / exit", AMBER), ("Re-ID", "HSV torso histogram", VIOLET),
        ("PID control", "steer + distance", GREEN)]
for i, (t, sub, col) in enumerate(flow):
    x = 0.6 + i * 1.5
    node(s, x, 1.55, 1.3, 1.0, t, sub, color=col, title_size=11.5, sub_size=9)
    if i < len(flow) - 1:
        arrow(s, x + 1.3, 2.05, x + 1.5, 2.05)
card(s, 0.6, 2.85, 7.3, 1.9, "Re-identification (tracker.py)", [
    "2-D H×S histogram (30×32 bins) of the upper torso stored on lock",
    "Candidates compared with Bhattacharyya distance",
    "Best match above the similarity threshold becomes the target again",
], accent=VIOLET, body_size=12)
card(s, 0.6, 4.9, 7.3, 1.9, "PID visual servoing (pid.py)", [
    "u(t) = Kp·e + Ki·Σe·Δt + Kd·Δe/Δt",
    "Steering error = (x_box − x_centre) / frame width",
    "Distance error = h_box / h_frame − 0.30 (target box height)",
], accent=GREEN, body_size=12)
placeholder(s, 8.2, 1.55, 4.55, 5.25, "Follow-me demo frame", "Target box + re-ID score + PID output")

# =====================================================================
# SLIDE 9 — FIRMWARE FSM
# =====================================================================
s = new_slide("Firmware — 3-Tier Priority State Machine",
              "Safety decisions that never wait for Linux",
              notes="Tier 1 always wins. Tier 2 reacts to sonar on the Arduino itself. Tier 3 executes the Pi's "
                    "movement commands. Values below come from protocol.h / the firmware.")
tiers = [
    ("TIER 1 — Emergency", RED, [
        "PI_STOP: PWM → 0 instantly, latch until PI_GO",
        "Watchdog: no valid packet within WATCHDOG_TIMEOUT_MS → motors off",
        "Movement commands rejected while latched"]),
    ("TIER 2 — Local reflex", AMBER, [
        "Trigger: either sonar < 20 cm",
        "Reverse 400 ms → turn away 500 ms → resume",
        "PI_AVOID hint overrides turn direction; 500 ms cooldown"]),
    ("TIER 3 — Navigation", GREEN, [
        "DRIVE <speed> <ms>, REV, TURN <deg> <speed>, STOP",
        "Hardware Timer0 PWM on D5 / D6 (0–255)",
        "Sent by the Pi 5 decision arbiter"]),
]
for i, (t, col, lines) in enumerate(tiers):
    y = 1.5 + i * 1.78
    w = 7.6 - i * 0.0
    card(s, 0.6, y, w, 1.6, t, lines, accent=col, body_size=12)
placeholder(s, 8.5, 1.5, 4.25, 5.3, "Serial monitor log", "e.g. [TELEM] AVOID triggered / [ACK] BLOCKED (PI_STOPPED)")

# =====================================================================
# SLIDE 10 — OUR BENCHMARKS
# =====================================================================
s = new_slide("Our Benchmark Results (Raspberry Pi 5)",
              "Fill with numbers measured on the car — cells marked __",
              notes="Run benchmark.py on the Pi 5 and copy the numbers here. Keep the screenshot as evidence.")
text(s, 0.6, 1.45, 6, 0.4, "Runtime comparison (imgsz 320, 100 frames)", 14, CYAN, bold=True)
table(s, 0.6, 1.9, 6.4, 2.2, [
    ["Format", "Avg latency (ms)", "Median (ms)", "FPS", "Det / frame"],
    ["PyTorch FP32", "__", "__", "__", "__"],
    ["ONNX FP32", "__", "__", "__", "__"],
    ["NCNN FP16", "__", "__", "__", "__"],
    ["OpenVINO FP16", "__", "__", "__", "__"],
], col_widths=[1.6, 1.4, 1.2, 0.9, 1.3], size=11, highlight_row=3)
text(s, 0.6, 4.35, 6, 0.4, "Resolution sweep (NCNN FP16)", 14, CYAN, bold=True)
table(s, 0.6, 4.8, 6.4, 1.7, [
    ["Input size", "Avg latency (ms)", "FPS", "Det / frame"],
    ["640 × 640", "__", "__", "__"],
    ["480 × 480", "__", "__", "__"],
    ["320 × 320", "__", "__", "__"],
], col_widths=[1.6, 1.8, 1.4, 1.6], size=11)
placeholder(s, 7.3, 1.45, 5.45, 2.6, "Terminal: benchmark comparison table", "python benchmark.py --video dashcam_test.mp4")
placeholder(s, 7.3, 4.25, 5.45, 2.55, "htop / vcgencmd measure_temp during run", "CPU load & temperature under NCNN")

# =====================================================================
# SLIDE 11 — COMPARISON WITH OTHER PLATFORMS
# =====================================================================
s = new_slide("Comparison with Similar Platforms",
              "Perception stack and obstacle handling (see report §2 for sources)",
              notes="Values are taken from the comparative benchmarking report, section 2, with citations there.")
table(s, 0.6, 1.5, 12.15, 4.6, [
    ["Platform", "Primary model", "Params", "NMS step", "Compute", "Obstacle sensing", "Pothole"],
    ["Our vehicle", "YOLOv10n (NCNN)", "2.3 M", "None", "Pi 5 CPU", "Camera + 2× sonar reflex", "Yes"],
    ["NVIDIA JetBot", "SSD-MobileNetV2 / ResNet-18", "4.3 M / 11.7 M", "Yes (SSD)", "Jetson Nano GPU", "Camera classifier", "No"],
    ["AWS DeepRacer", "PPO policy CNN", "small", "N/A", "Intel Atom", "Camera / LiDAR (RL)", "No"],
    ["Duckietown DB21", "Classical lane pipeline", "—", "N/A", "Jetson Nano", "AprilTags + camera", "No"],
    ["OpenBot", "MobileNet-SSD / CNN", "~4 M", "Yes", "Smartphone", "1× sonar + camera", "No"],
    ["DonkeyCar", "Behavioural-cloning CNN", "~1.6 M", "N/A", "Pi 4B", "None (track racing)", "No"],
    ["SunFounder PiCar-X", "TFLite demos", "—", "Yes", "Pi 4B", "1× sonar, grayscale", "No"],
    ["Hiwonder MentorPi", "YOLOv5s", "7.2 M", "Yes", "Pi 5", "LiDAR + sonar", "No"],
], col_widths=[1.75, 2.35, 1.35, 1.1, 1.6, 2.55, 1.45], size=11, highlight_row=1)
text(s, 0.6, 6.3, 12, 0.6,
     "Verify every row against report §2 before presenting; keep citations on the report slide/handout.",
     11, MUTED, italic=True)

# =====================================================================
# SLIDE 12 — LIVE DEMO
# =====================================================================
s = new_slide("Live Demo — YOLOv10n on Any Video",
              "test_yolov10_video.py: HUD with latency, FPS and object counts",
              notes="Run the script live. Use 'p' to pause on an interesting frame and explain the boxes.")
card(s, 0.6, 1.5, 5.6, 3.3, "Commands", [
    "python test_yolov10_video.py --video pothole_test_2.mp4",
    "python test_yolov10_video.py --video <any_video.mp4>",
    "python test_yolov10_video.py --webcam 0",
    "add  --save demo.mp4  to record",
], accent=CYAN, body_size=12)
card(s, 0.6, 5.0, 5.6, 1.8, "Controls", [
    "p  pause / resume     s  screenshot     q  quit",
], accent=GREEN, body_size=12)
placeholder(s, 6.5, 1.5, 6.25, 5.3, "Demo HUD frame", "Bounding boxes + latency / FPS overlay")

# =====================================================================
# SLIDE 13 — ROADMAP
# =====================================================================
s = new_slide("Roadmap — Next Integrations",
              "Three capabilities, nothing else",
              notes="Pothole integration on the live car, lane following, and traffic sign recognition.")
card(s, 0.6, 1.5, 3.9, 4.6, "01  Pothole avoidance", [
    "Run RDD2022 YOLOv10n + CV fallback in the live loop",
    "Validate corridor + swerve decisions on real surfaces",
    "Tune POTHOLE_WEIGHT and thresholds",
], accent=AMBER)
card(s, 4.72, 1.5, 3.9, 4.6, "02  Lane following", [
    "Multi-ROI lane detector in LAB colour space",
    "Heading error θ = arctan(dx / dy)",
    "Feed θ into the existing PID controller",
], accent=CYAN)
card(s, 8.84, 1.5, 3.9, 4.6, "03  Traffic signs & signals", [
    "6 classes: go, right, park, red, green, crosswalk",
    "5-frame debounce + upper-ROI gating",
    "Red = stop latch, green = release, crosswalk = slow",
], accent=GREEN)
text(s, 0.6, 6.3, 12, 0.5, "Status update: ____________________", 13, MUTED)

# =====================================================================
# SLIDE 14 — THANK YOU
# =====================================================================
s = new_slide(notes="Open the floor for questions.")
text(s, 0, 2.5, 13.333, 1.0, "Thank You", 54, TEXT, bold=True, align=PP_ALIGN.CENTER)
text(s, 0, 3.6, 13.333, 0.5, "Questions?", 22, CYAN, align=PP_ALIGN.CENTER)
text(s, 0, 4.5, 13.333, 0.5, "Krishna Menon  •  Aditya Rana  •  Parth Surve", 16, MUTED, align=PP_ALIGN.CENTER)

# ---------------- SAVE ----------------
os.makedirs(OUT_DIR, exist_ok=True)
out = os.path.join(OUT_DIR, OUT_FILE)
prs.save(out)
print(f"Saved: {out}  ({len(prs.slides)} slides)")
