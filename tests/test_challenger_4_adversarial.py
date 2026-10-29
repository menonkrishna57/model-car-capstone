r"""
test_challenger_4_adversarial.py
================================
Adversarial Stress Test Suite for Gate 2:
Anti-AI Prose & Data Integrity Challenger

Target Document: Autonomous_RC_Car_Capstone_Report_Final.docx
Python Env: model/.venv/Scripts/python.exe

Checks:
1. Banned buzzword audit (17 words across every paragraph and table cell, assert 0 Crucial/crucially)
2. Dash audit (0 em-dashes, 0 en-dashes, 0 spaced double hyphens)
3. Markdown & Hash audit (0 markdown artifacts, 0 literal '#' in any paragraph)
4. Table cleanliness (all 31 tables, 0 empty or whitespace-only cells, including Table 0)
5. Image audit (all 6 figures from model/paper_figures/ embedded cleanly)
"""

import os
import sys
import re
import zipfile
import hashlib
import xml.etree.ElementTree as ET
import docx

sys.stdout.reconfigure(encoding='utf-8')

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

BANNED_BUZZWORDS = [
    "crucial", "crucially", "delve", "testament", "tapestry", "landscape", "pivotal",
    "robust", "showcase", "underscore", "foster", "interplay", "bolster",
    "beacon", "catalyst", "panacea", "plethora"
]

def hash_file_bytes(b):
    return hashlib.sha256(b).hexdigest()

def audit_banned_buzzwords(doc, doc_xml):
    print("=" * 70)
    print("AUDIT 1: BANNED BUZZWORDS AUDIT")
    print("=" * 70)
    
    findings = []
    
    # 1. Paragraph scan
    for p_idx, p in enumerate(doc.paragraphs):
        text = p.text
        for word in BANNED_BUZZWORDS:
            # Word boundary search (case-insensitive)
            matches = list(re.finditer(rf'\b{re.escape(word)}\b', text, re.IGNORECASE))
            for m in matches:
                findings.append({
                    "location": f"Paragraph {p_idx}",
                    "word": word,
                    "matched": m.group(0),
                    "context": text[max(0, m.start()-30):min(len(text), m.end()+30)]
                })
                
    # 2. Table cell scan
    for t_idx, tbl in enumerate(doc.tables):
        for r_idx, row in enumerate(tbl.rows):
            for c_idx, cell in enumerate(row.cells):
                cell_text = cell.text
                for word in BANNED_BUZZWORDS:
                    matches = list(re.finditer(rf'\b{re.escape(word)}\b', cell_text, re.IGNORECASE))
                    for m in matches:
                        findings.append({
                            "location": f"Table {t_idx} Row {r_idx} Col {c_idx}",
                            "word": word,
                            "matched": m.group(0),
                            "context": cell_text[max(0, m.start()-30):min(len(cell_text), m.end()+30)]
                        })
                        
    # 3. Header & Footer scan
    for s_idx, sec in enumerate(doc.sections):
        for h_p_idx, p in enumerate(sec.header.paragraphs):
            for word in BANNED_BUZZWORDS:
                matches = list(re.finditer(rf'\b{re.escape(word)}\b', p.text, re.IGNORECASE))
                for m in matches:
                    findings.append({
                        "location": f"Section {s_idx} Header P{h_p_idx}",
                        "word": word,
                        "matched": m.group(0),
                        "context": p.text
                    })
        for f_p_idx, p in enumerate(sec.footer.paragraphs):
            for word in BANNED_BUZZWORDS:
                matches = list(re.finditer(rf'\b{re.escape(word)}\b', p.text, re.IGNORECASE))
                for m in matches:
                    findings.append({
                        "location": f"Section {s_idx} Footer P{f_p_idx}",
                        "word": word,
                        "matched": m.group(0),
                        "context": p.text
                    })

    # Specifically check for crucial / Crucially count
    crucial_matches = [f for f in findings if f["word"] in ("crucial", "crucially")]
    
    print(f"Total Banned Word Matches: {len(findings)}")
    print(f"Occurrences of 'crucial' / 'Crucially': {len(crucial_matches)}")
    
    if findings:
        for f in findings:
            print(f"  [VIOLATION] {f['location']}: word='{f['word']}' matched='{f['matched']}' -> '...{f['context']}...'")
    else:
        print("  [SUCCESS] Exactly 0 occurrences of all 17 banned words across all paragraphs, table cells, headers, and footers.")
        
    return len(findings) == 0, findings


def audit_dashes(doc, doc_xml):
    print("\n" + "=" * 70)
    print("AUDIT 2: DASH AUDIT")
    print("=" * 70)
    
    em_dash_violations = []
    en_dash_violations = []
    double_hyphen_violations = []
    
    # 1. Paragraphs
    for p_idx, p in enumerate(doc.paragraphs):
        t = p.text
        if '\u2014' in t: # Em-dash —
            em_dash_violations.append((f"Paragraph {p_idx}", t))
        if '\u2013' in t: # En-dash –
            # Note: check if preformatted code
            em_type = "Preformatted" if (p.style.name and 'Code' in p.style.name) else "Prose"
            en_dash_violations.append((f"Paragraph {p_idx} ({em_type})", t))
        if ' -- ' in t:
            double_hyphen_violations.append((f"Paragraph {p_idx}", t))
            
    # 2. Table cells
    for t_idx, tbl in enumerate(doc.tables):
        for r_idx, row in enumerate(tbl.rows):
            for c_idx, cell in enumerate(row.cells):
                t = cell.text
                if '\u2014' in t:
                    em_dash_violations.append((f"Table {t_idx} R{r_idx}C{c_idx}", t))
                if '\u2013' in t:
                    en_dash_violations.append((f"Table {t_idx} R{r_idx}C{c_idx}", t))
                if ' -- ' in t:
                    double_hyphen_violations.append((f"Table {t_idx} R{r_idx}C{c_idx}", t))

    # 3. Headers / Footers
    for s_idx, sec in enumerate(doc.sections):
        for p in sec.header.paragraphs:
            if '\u2014' in p.text:
                em_dash_violations.append((f"Section {s_idx} Header", p.text))
            if '\u2013' in p.text:
                en_dash_violations.append((f"Section {s_idx} Header", p.text))
            if ' -- ' in p.text:
                double_hyphen_violations.append((f"Section {s_idx} Header", p.text))
        for p in sec.footer.paragraphs:
            if '\u2014' in p.text:
                em_dash_violations.append((f"Section {s_idx} Footer", p.text))
            if '\u2013' in p.text:
                en_dash_violations.append((f"Section {s_idx} Footer", p.text))
            if ' -- ' in p.text:
                double_hyphen_violations.append((f"Section {s_idx} Footer", p.text))

    print(f"Em-dashes (U+2014): {len(em_dash_violations)}")
    for loc, txt in em_dash_violations:
        print(f"  [VIOLATION] {loc}: {repr(txt[:60])}")
        
    print(f"En-dashes (U+2013): {len(en_dash_violations)}")
    for loc, txt in en_dash_violations:
        print(f"  [VIOLATION] {loc}: {repr(txt[:60])}")
        
    print(f"Spaced double hyphens (' -- '): {len(double_hyphen_violations)}")
    for loc, txt in double_hyphen_violations:
        print(f"  [VIOLATION] {loc}: {repr(txt[:60])}")

    passed = (len(em_dash_violations) == 0 and len(en_dash_violations) == 0 and len(double_hyphen_violations) == 0)
    if passed:
        print("  [SUCCESS] Exactly 0 em-dashes, 0 en-dashes, and 0 spaced double hyphens detected.")
        
    return passed, {
        "em_dashes": em_dash_violations,
        "en_dashes": en_dash_violations,
        "double_hyphens": double_hyphen_violations
    }


def audit_markdown_and_hashes(doc):
    print("\n" + "=" * 70)
    print("AUDIT 3: MARKDOWN & HASH AUDIT")
    print("=" * 70)
    
    hash_violations = []
    markdown_violations = []
    
    for p_idx, p in enumerate(doc.paragraphs):
        t = p.text
        # Literal '#' check in ANY paragraph
        if '#' in t:
            hash_violations.append((p_idx, t))
            
        stripped = t.strip()
        # Markdown heading hashes
        if re.match(r'^\s*#{1,6}\s+', stripped):
            markdown_violations.append((f"P{p_idx}", "Markdown heading hash", stripped[:50]))
        # Blockquote
        if stripped.startswith('>'):
            markdown_violations.append((f"P{p_idx}", "Blockquote chevron", stripped[:50]))
        # Double asterisks (bold)
        if '**' in t:
            markdown_violations.append((f"P{p_idx}", "Unparsed ** bold", stripped[:50]))
        # Inline backticks outside code style
        if '`' in t:
            is_code_style = p.style.name and ('Code' in p.style.name or 'Preformatted' in p.style.name)
            if not is_code_style:
                markdown_violations.append((f"P{p_idx}", "Inline backtick", stripped[:50]))
        # Markdown links [text](url)
        if re.search(r'\[([^\]]+)\]\((https?://[^\)]+)\)', t):
            markdown_violations.append((f"P{p_idx}", "Markdown link", stripped[:50]))
        # Horizontal rules
        if re.match(r'^([-*_]){3,}$', stripped):
            markdown_violations.append((f"P{p_idx}", "Markdown horizontal rule", stripped[:50]))
        # Raw bullet list marker
        if re.match(r'^\s*[\*\+]\s+', stripped):
            markdown_violations.append((f"P{p_idx}", "Raw bullet marker (* or +)", stripped[:50]))

    # Also check table cells for '#' and markdown
    for t_idx, tbl in enumerate(doc.tables):
        for r_idx, row in enumerate(tbl.rows):
            for c_idx, cell in enumerate(row.cells):
                ct = cell.text
                if '#' in ct:
                    hash_violations.append((f"Table {t_idx} R{r_idx}C{c_idx}", ct))
                if '**' in ct:
                    markdown_violations.append((f"Table {t_idx} R{r_idx}C{c_idx}", "Unparsed ** bold", ct[:40]))
                if '`' in ct:
                    markdown_violations.append((f"Table {t_idx} R{r_idx}C{c_idx}", "Inline backtick", ct[:40]))
                if re.match(r'^\s*#{1,6}\s+', ct.strip()):
                    markdown_violations.append((f"Table {t_idx} R{r_idx}C{c_idx}", "Heading hash", ct[:40]))

    print(f"Literal '#' violations: {len(hash_violations)}")
    for loc, txt in hash_violations:
        print(f"  [VIOLATION] {loc}: {repr(txt[:60])}")
        
    print(f"Markdown artifact violations: {len(markdown_violations)}")
    for loc, kind, txt in markdown_violations:
        print(f"  [VIOLATION] {loc} ({kind}): {repr(txt)}")

    passed = (len(hash_violations) == 0 and len(markdown_violations) == 0)
    if passed:
        print("  [SUCCESS] 0 literal '#' characters and 0 markdown artifacts found.")
        
    return passed, {"hashes": hash_violations, "markdown": markdown_violations}


def audit_table_cleanliness(doc):
    print("\n" + "=" * 70)
    print("AUDIT 4: TABLE CLEANLINESS AUDIT (ALL 31 TABLES)")
    print("=" * 70)
    
    total_tables = len(doc.tables)
    print(f"Total Tables in Document: {total_tables}")
    
    empty_cells = []
    total_cells = 0
    table_stats = []
    
    for t_idx, tbl in enumerate(doc.tables):
        num_rows = len(tbl.rows)
        num_cols = len(tbl.columns) if tbl.rows else 0
        tbl_cells = 0
        tbl_empty = 0
        
        # Check first row text to identify table
        first_row_text = " | ".join(c.text.strip().replace('\n', ' ') for c in tbl.rows[0].cells[:4]) if tbl.rows else ""
        
        for r_idx, row in enumerate(tbl.rows):
            for c_idx, cell in enumerate(row.cells):
                tbl_cells += 1
                total_cells += 1
                raw_text = cell.text
                if not raw_text or not raw_text.strip():
                    empty_cells.append({
                        "table_idx": t_idx,
                        "row_idx": r_idx,
                        "col_idx": c_idx,
                        "content": repr(raw_text)
                    })
                    tbl_empty += 1
                    
        table_stats.append({
            "idx": t_idx,
            "rows": num_rows,
            "cols": num_cols,
            "cells": tbl_cells,
            "empty": tbl_empty,
            "preview": first_row_text[:60]
        })

    for s in table_stats:
        status = "CLEAN" if s["empty"] == 0 else f"VIOLATIONS ({s['empty']} empty)"
        print(f"  Table {s['idx']:2d}: {s['rows']:2d} rows x {s['cols']:2d} cols ({s['cells']:3d} cells) [{status}] - First row: {s['preview']}")

    print(f"\nTotal Cells Audited across {total_tables} tables: {total_cells}")
    print(f"Total Empty / Whitespace-Only Cells: {len(empty_cells)}")
    
    if empty_cells:
        for ec in empty_cells:
            print(f"  [VIOLATION] Table {ec['table_idx']} Row {ec['row_idx']} Col {ec['col_idx']}: {ec['content']}")
    else:
        print("  [SUCCESS] Exactly 0 empty or whitespace-only cells across all tables.")
        
    # Check if total tables is exactly 31
    table_count_ok = (total_tables == 31)
    if not table_count_ok:
        print(f"  [NOTE] Total tables is {total_tables}, requirement states 31 tables.")
        
    passed = (len(empty_cells) == 0 and total_tables == 31)
    return passed, {"total_tables": total_tables, "total_cells": total_cells, "empty_cells": empty_cells, "stats": table_stats}


def audit_images(doc, docx_path, figures_dir):
    print("\n" + "=" * 70)
    print("AUDIT 5: IMAGE AUDIT (6 FIGURES FROM model/paper_figures/)")
    print("=" * 70)
    
    expected_files = [
        "fig1_system_topology.png",
        "fig2_power_distribution.png",
        "fig3_chassis_layout.png",
        "fig4_state_machine.png",
        "fig5_perception_pipeline.png",
        "fig6_sensor_physics_and_corridor.png"
    ]
    
    # 1. Check source figures in figures_dir
    source_hashes = {}
    source_sizes = {}
    for fname in expected_files:
        fpath = os.path.join(figures_dir, fname)
        if not os.path.exists(fpath):
            print(f"  [FAIL] Source figure missing on disk: {fpath}")
            return False, {}
        with open(fpath, 'rb') as f:
            data = f.read()
            source_hashes[fname] = hash_file_bytes(data)
            source_sizes[fname] = len(data)
        print(f"  Source {fname}: size={source_sizes[fname]} bytes, sha256={source_hashes[fname][:16]}...")
        
    # 2. Inspect embedded media inside docx
    embedded_media = {}
    with zipfile.ZipFile(docx_path, 'r') as z:
        media_names = [n for n in z.namelist() if n.startswith('word/media/')]
        for mn in sorted(media_names):
            data = z.read(mn)
            h = hash_file_bytes(data)
            embedded_media[mn] = {
                "size": len(data),
                "sha256": h
            }
            print(f"  Embedded {mn}: size={len(data)} bytes, sha256={h[:16]}...")

        # 3. Inspect relationships
        rels_xml = z.read('word/_rels/document.xml.rels').decode('utf-8')
        rels_root = ET.fromstring(rels_xml)
        image_rels = {}
        for r in rels_root.findall(f'.//{{{R_NS}}}Relationship'):
            rtype = r.attrib.get('Type', '')
            if rtype.endswith('/image'):
                rid = r.attrib.get('Id')
                target = r.attrib.get('Target')
                image_rels[rid] = target

        # 4. Inspect document.xml drawing references and captions
        doc_xml = z.read('word/document.xml').decode('utf-8')
        doc_root = ET.fromstring(doc_xml)
        body = doc_root.find(f'{{{W_NS}}}body')
        
        drawing_captions = []
        for idx, el in enumerate(list(body)):
            # Check if element has drawing
            drawings = el.findall(f'.//{{{WP_NS}}}inline') + el.findall(f'.//{{{WP_NS}}}anchor')
            if drawings:
                # Find the blip r:embed
                blips = el.findall(f'.//{{{A_NS}}}blip')
                embed_rids = [b.attrib.get(f'{{{R_NS}}}embed') for b in blips if f'{{{R_NS}}}embed' in b.attrib]
                # Find caption in following paragraph
                next_el = body[idx + 1] if idx + 1 < len(body) else None
                caption_text = ''.join(next_el.itertext()).strip() if next_el is not None else ""
                drawing_captions.append({
                    "body_idx": idx,
                    "embed_rids": embed_rids,
                    "targets": [image_rels.get(r, "UNKNOWN") for r in embed_rids],
                    "caption": caption_text
                })

    print(f"\nTotal media files in docx: {len(embedded_media)}")
    print(f"Total drawings in document body: {len(drawing_captions)}")
    
    for i, dc in enumerate(drawing_captions):
        print(f"  Drawing {i+1}: targets={dc['targets']}, caption='{dc['caption'][:60]}'")

    # Verify matching of hashes
    source_matched = 0
    for sname, shash in source_hashes.items():
        found = False
        for em_name, em_info in embedded_media.items():
            if em_info["sha256"] == shash:
                found = True
                print(f"  [MATCH] {sname} corresponds exactly to {em_name} (sha256 verified)")
                break
        if found:
            source_matched += 1
        else:
            print(f"  [MISMATCH] {sname} hash not matched in embedded docx media!")

    all_captions_below = all(
        re.match(r'^Figure\s+\d+(\.\d+)?:\s+', dc['caption']) for dc in drawing_captions
    )
    
    passed = (
        len(drawing_captions) == 6 and
        len(embedded_media) >= 6 and
        source_matched == 6 and
        all_captions_below
    )
    
    if passed:
        print("  [SUCCESS] All 6 figures are cleanly embedded with exact SHA-256 match and Figure captions below.")
    else:
        print(f"  [FAIL] Image audit failed: drawings={len(drawing_captions)}, media={len(embedded_media)}, matched={source_matched}, captions_below={all_captions_below}")
        
    return passed, {
        "drawings": drawing_captions,
        "media": embedded_media,
        "source_matched": source_matched,
        "captions_below": all_captions_below
    }


def main():
    repo_root = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
    docx_path = os.path.join(repo_root, "Autonomous_RC_Car_Capstone_Report_Final.docx")
    figures_dir = os.path.join(repo_root, "model", "paper_figures")
    
    print("=" * 80)
    print("GATE 2 ADVERSARIAL STRESS TEST SUITE: ANTI-AI PROSE & DATA INTEGRITY")
    print(f"Document Path: {docx_path}")
    print(f"Figures Dir:   {figures_dir}")
    print("=" * 80)
    
    if not os.path.exists(docx_path):
        print(f"FATAL: Document not found at {docx_path}")
        sys.exit(1)
        
    doc = docx.Document(docx_path)
    with zipfile.ZipFile(docx_path, 'r') as z:
        doc_xml = z.read('word/document.xml').decode('utf-8')
        
    p1, r1 = audit_banned_buzzwords(doc, doc_xml)
    p2, r2 = audit_dashes(doc, doc_xml)
    p3, r3 = audit_markdown_and_hashes(doc)
    p4, r4 = audit_table_cleanliness(doc)
    p5, r5 = audit_images(doc, docx_path, figures_dir)
    
    print("\n" + "=" * 80)
    print("FINAL SUMMARY OF GATE 2 AUDIT")
    print("=" * 80)
    print(f"1. Banned Buzzwords (17 banned words):        {'PASS' if p1 else 'FAIL'} (violations: {len(r1)})")
    print(f"2. Dashes (0 em, 0 en, 0 ' -- '):             {'PASS' if p2 else 'FAIL'} (em: {len(r2['em_dashes'])}, en: {len(r2['en_dashes'])}, double: {len(r2['double_hyphens'])})")
    print(f"3. Markdown & Hash (0 artifacts, 0 literal '#'): {'PASS' if p3 else 'FAIL'} (hashes: {len(r3['hashes'])}, md: {len(r3['markdown'])})")
    print(f"4. Table Cleanliness (31 tables, 0 empty):    {'PASS' if p4 else 'FAIL'} (tables: {r4['total_tables']}, empty cells: {len(r4['empty_cells'])})")
    print(f"5. Image Audit (6 figures cleanly embedded):  {'PASS' if p5 else 'FAIL'} (drawings: {len(r5.get('drawings', []))}, matched: {r5.get('source_matched', 0)})")
    
    all_passed = p1 and p2 and p3 and p4 and p5
    verdict = "APPROVE" if all_passed else "REQUEST_CHANGES"
    print("\nOVERALL VERDICT: " + verdict)
    print("=" * 80)
    
    return 0 if all_passed else 1


if __name__ == "__main__":
    sys.exit(main())
